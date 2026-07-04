
#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Apply NSG rules from an Excel file in Azure Cloud Shell.

Usage:
  python3 apply_nsg_rules.py --file Rules.xlsx

Key options:
  --what-if                 : Validate and show actions; do not call Azure.
  --stop-on-error           : Stop at the first failing row (default: continue).
  --sheet-name SHEET        : If your Excel has multiple sheets (default: first sheet).
  --log-file FILE.csv       : Write per-row results to a CSV log.
  --create-missing-nsgs     : Auto-create NSGs that are referenced but not found.
"""

import argparse
import sys
import subprocess
import shlex
from typing import List, Union, Optional
import csv
from datetime import datetime

# --- Robust import for openpyxl in Cloud Shell ---
try:
    import openpyxl  # type: ignore
except ImportError:
    print("Installing dependency: openpyxl ...")
    subprocess.run(
        [sys.executable, "-m", "pip", "install", "--user", "--upgrade", "openpyxl", "et-xmlfile"],
        check=True
    )
    try:
        import site
        user_site = site.getusersitepackages()
        if user_site and user_site not in sys.path:
            sys.path.append(user_site)
    except Exception:
        pass
    try:
        from pathlib import Path
        p = Path.home() / f".local/lib/python{sys.version_info[0]}.{sys.version_info[1]}/site-packages"
        if str(p) not in sys.path:
            sys.path.insert(0, str(p))
    except Exception:
        pass
    import openpyxl  # type: ignore


def title_or_star(val: str) -> str:
    if val is None:
        return '*'
    s = str(val).strip()
    if s == '' or s == '*':
        return '*'
    return s.title()


def normalize_any(val: Union[str, None]) -> str:
    if val is None:
        return '*'
    s = str(val).strip()
    return s if s else '*'


def split_multi(value: str) -> List[str]:
    parts = [p.strip() for p in str(value).split(',')]
    return [p for p in parts if p]


def parse_ports(raw: Union[str, None]) -> Union[str, List[str]]:
    s = normalize_any(raw)
    if s == '*':
        return '*'
    tokens = split_multi(s)
    out = []
    for t in tokens:
        if '-' in t:
            a, b = t.split('-', 1)
            if not (a.isdigit() and b.isdigit()):
                raise ValueError(f"Invalid port range '{t}'")
            lo, hi = int(a), int(b)
            if lo < 1 or hi > 65535 or lo > hi:
                raise ValueError(f"Port range out of bounds '{t}'")
            out.append(f"{lo}-{hi}")
        else:
            if not t.isdigit():
                raise ValueError(f"Invalid port '{t}'")
            p = int(t)
            if p < 1 or p > 65535:
                raise ValueError(f"Port out of bounds '{t}'")
            out.append(str(p))
    return out


def parse_priority(raw: Union[str, int]) -> int:
    if raw is None or str(raw).strip() == '':
        raise ValueError("Priority is required")
    try:
        pr = int(raw)
    except Exception:
        raise ValueError(f"Priority must be an integer, got '{raw}'")
    if pr < 100 or pr > 4096:
        raise ValueError("Priority must be between 100 and 4096")
    return pr


def norm_direction(raw: str) -> str:
    s = title_or_star(raw)
    if s not in ('Inbound', 'Outbound'):
        raise ValueError(f"Direction must be Inbound or Outbound, got '{raw}'")
    return s


def norm_protocol(raw: str) -> str:
    s = title_or_star(raw)
    if s not in ('Tcp', 'Udp', 'Icmp', 'Ah', 'Esp', '*'):
        raise ValueError(f"Protocol must be Tcp/Udp/Icmp/Ah/Esp or '*', got '{raw}'")
    return s


def norm_action(raw: Union[str, None]) -> str:
    """Map Action column to Azure 'access' flag."""
    if raw is None or str(raw).strip() == '':
        raise ValueError("Action is required (Allow or Deny)")
    s = str(raw).strip().title()
    if s not in ('Allow', 'Deny'):
        raise ValueError(f"Action must be Allow or Deny, got '{raw}'")
    return s


def address_args(flag_single: str, flag_multi: str, raw: Union[str, None]) -> List[str]:
    s = normalize_any(raw)
    if s == '*':
        return [flag_single, '*']
    values = split_multi(s)
    if len(values) == 1:
        return [flag_single, values[0]]
    return [flag_multi] + values


def port_args(flag_single: str, flag_multi: str, parsed: Union[str, List[str]]) -> List[str]:
    if parsed == '*':
        return [flag_single, '*']
    vals = list(parsed)
    if len(vals) == 1:
        return [flag_single, vals[0]]
    return [flag_multi] + vals


def az(*args) -> subprocess.CompletedProcess:
    cmd = ['az'] + list(args)
    return subprocess.run(cmd, check=True, capture_output=True, text=True)


def az_exists_rule(subscription: str, rg: str, nsg: str, rule: str) -> bool:
    try:
        az('network', 'nsg', 'rule', 'show',
           '--subscription', subscription,
           '--resource-group', rg,
           '--nsg-name', nsg,
           '--name', rule)
        return True
    except subprocess.CalledProcessError:
        return False


def az_exists_nsg(subscription: str, rg: str, nsg: str) -> bool:
    try:
        az('network', 'nsg', 'show',
           '--subscription', subscription,
           '--resource-group', rg,
           '--name', nsg)
        return True
    except subprocess.CalledProcessError:
        return False


def az_create_nsg(subscription: str, rg: str, nsg: str) -> None:
    az('network', 'nsg', 'create',
       '--subscription', subscription,
       '--resource-group', rg,
       '--name', nsg)


def build_rule_command(
    create: bool,
    subscription: str,
    rg: str,
    nsg: str,
    name: str,
    priority: int,
    direction: str,
    action: str,
    protocol: str,
    src_addr_raw: Union[str, None],
    dst_addr_raw: Union[str, None],
    src_ports_raw: Union[str, None],
    dst_ports_raw: Union[str, None],
    description: Union[str, None]
) -> List[str]:
    cmd = [
        'create' if create else 'update',
        '--subscription', subscription,
        '--resource-group', rg,
        '--nsg-name', nsg,
        '--name', name,
        '--priority', str(priority),
        '--direction', direction,
        '--access', action,           # <-- Action mapped to --access
        '--protocol', protocol
    ]
    cmd += address_args('--source-address-prefix', '--source-address-prefixes', src_addr_raw)
    cmd += address_args('--destination-address-prefix', '--destination-address-prefixes', dst_addr_raw)
    src_ports = parse_ports(src_ports_raw)
    dst_ports = parse_ports(dst_ports_raw)
    cmd += port_args('--source-port-range', '--source-port-ranges', src_ports)
    cmd += port_args('--destination-port-range', '--destination-port-ranges', dst_ports)
    if description and str(description).strip():
        cmd += ['--description', str(description).strip()]
    return cmd


def write_log_header(csv_path: str):
    with open(csv_path, 'w', newline='') as f:
        writer = csv.writer(f)
        writer.writerow([
            'timestamp',
            'row_number',
            'subscription',
            'resource_group',
            'nsg_name',
            'rule_name',
            'action_type',   # CREATE / UPDATE / WHAT-IF / FAILED / SKIPPED / CREATE-NSG
            'status',        # success / failed
            'message',
            'az_command'
        ])


def append_log(
    csv_path: Optional[str],
    row_number: int,
    subscription: str,
    rg: str,
    nsg: str,
    rule: str,
    action_type: str,
    status: str,
    message: str,
    az_command: str
):
    if not csv_path:
        return
    with open(csv_path, 'a', newline='') as f:
        writer = csv.writer(f)
        writer.writerow([
            datetime.utcnow().isoformat() + 'Z',
            row_number,
            subscription,
            rg,
            nsg,
            rule,
            action_type,
            status,
            message,
            az_command
        ])


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--file', '-f', required=True, help='Path to Rules.xlsx')
    parser.add_argument('--sheet-name', default=None, help='Excel sheet name (default: first sheet)')
    parser.add_argument('--what-if', dest='what_if', action='store_true', help='Validate and show actions, but do not call Azure')
    parser.add_argument('--stop-on-error', action='store_true', help='Stop on first error (default: continue)')
    parser.add_argument('--log-file', default=None, help='CSV file to write per-row results')
    parser.add_argument('--create-missing-nsgs', action='store_true', help='Auto-create NSGs that are referenced but not found')
    args = parser.parse_args()

    if args.log_file:
        write_log_header(args.log_file)

    # Open workbook
    wb = openpyxl.load_workbook(args.file, read_only=True, data_only=True)
    ws = wb[args.sheet_name] if args.sheet_name else wb.worksheets[0]

    # Read header mapping
    headers = [str(c.value).strip() if c.value is not None else '' for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=False))]
    header_map = {h: i for i, h in enumerate(headers)}

    required = [
        'NSGName', 'RuleName', 'Description', 'Direction', 'Priority', 'Protocol',
        'SourceAddress', 'DestinationAddress', 'SourcePort', 'DestinationPort',
        'Subscription', 'ResourceGroup', 'Action'
    ]
    for h in required:
        if h not in header_map:
            msg = f"Missing required column '{h}' in header."
            print(f"ERROR: {msg}")
            append_log(args.log_file, 0, '', '', '', '', 'INIT', 'failed', msg, '')
            sys.exit(1)

    total = 0
    created = 0
    updated = 0
    skipped = 0
    failed = 0

    print("Starting NSG rule application...\n")

    for row in ws.iter_rows(min_row=2, values_only=True):
        total += 1
        subscription = rg = nsg = name = ''  # for safe logging in except
        try:
            nsg = str(row[header_map['NSGName']]).strip()
            name = str(row[header_map['RuleName']]).strip()
            description = row[header_map['Description']]
            direction = norm_direction(row[header_map['Direction']])
            priority = parse_priority(row[header_map['Priority']])
            protocol = norm_protocol(row[header_map['Protocol']])
            src_addr = row[header_map['SourceAddress']]
            dst_addr = row[header_map['DestinationAddress']]
            src_port = row[header_map['SourcePort']]
            dst_port = row[header_map['DestinationPort']]
            subscription = str(row[header_map['Subscription']]).strip()
            rg = str(row[header_map['ResourceGroup']]).strip()
            action = norm_action(row[header_map['Action']])  # <-- NEW: required Action

            if not nsg or not name or not subscription or not rg:
                raise ValueError("NSGName, RuleName, Subscription, and ResourceGroup are required fields")

            # Ensure NSG exists or create it
            if not args.what_if:
                if not az_exists_nsg(subscription, rg, nsg):
                    if args.create_missing_nsgs:
                        print(f"[INFO] NSG '{nsg}' not found. Creating...")
                        az_create_nsg(subscription, rg, nsg)
                        append_log(args.log_file, total, subscription, rg, nsg, name,
                                   'CREATE-NSG', 'success', 'NSG created',
                                   f"az network nsg create --subscription {subscription} --resource-group {rg} --name {nsg}")
                    else:
                        raise ValueError(f"NSG '{nsg}' not found in RG '{rg}' (subscription '{subscription}')")

            exists = False if args.what_if else az_exists_rule(subscription, rg, nsg, name)

            cmd = build_rule_command(
                create=(not exists),
                subscription=subscription,
                rg=rg,
                nsg=nsg,
                name=name,
                priority=priority,
                direction=direction,
                action=action,
                protocol=protocol,
                src_addr_raw=src_addr,
                dst_addr_raw=dst_addr,
                src_ports_raw=src_port,
                dst_ports_raw=dst_port,
                description=description
            )

            az_print = 'az network nsg rule ' + ' '.join(shlex.quote(c) for c in cmd)

            if args.what_if:
                print(f"[WHAT-IF] {'CREATE' if not exists else 'UPDATE'}: {az_print}")
                skipped += 1
                append_log(args.log_file, total, subscription, rg, nsg, name,
                           'WHAT-IF', 'success', 'No changes applied', az_print)
                continue

            az('network', 'nsg', 'rule', *cmd)
            if exists:
                print(f"[UPDATED] {rg}/{nsg}/{name}")
                updated += 1
                append_log(args.log_file, total, subscription, rg, nsg, name,
                           'UPDATE', 'success', 'Rule updated', az_print)
            else:
                print(f"[CREATED] {rg}/{nsg}/{name}")
                created += 1
                append_log(args.log_file, total, subscription, rg, nsg, name,
                           'CREATE', 'success', 'Rule created', az_print)

        except Exception as e:
            msg = str(e)
            print(f"[FAILED] Row {total}: {msg}")
            failed += 1
            append_log(args.log_file, total, subscription, rg, nsg, name,
                       'FAILED', 'failed', msg, az_print if 'az_print' in locals() else '')
            if args.stop_on_error:
                break

    print("\n=== Summary ===")
    print(f"Total rows:   {total}")
    print(f"Created:      {created}")
    print(f"Updated:      {updated}")
    print(f"Skipped:      {skipped}")
    print(f"Failed:       {failed}")

    if failed > 0:
        sys.exit(2)


if __name__ == "__main__":
    main()
